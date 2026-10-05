"""Unit tests for document parsers: text, markdown, json, csv, pdf, docx, limits, and errors."""

import json
from pathlib import Path
from unittest.mock import MagicMock, patch
import zipfile

import docx
from pypdf import PdfWriter
import pytest

from modules.ingestion.parsers import (
    DOCX_EXPANDED_LIMIT,
    PDF_PAGE_LIMIT,
    ParsedDocument,
    _parse_csv,
    _parse_docx,
    _parse_pdf,
    _read_utf8,
    parse_file,
    parse_file_bounded,
)


class TestTextAndMarkdownParsers:
    """Test plain text and markdown parsing, including UTF-8 BOM handling."""

    def test_read_utf8_plain_and_with_bom(self, tmp_path: Path) -> None:
        plain_file = tmp_path / "plain.txt"
        plain_file.write_text("Hello, world!", encoding="utf-8")
        assert _read_utf8(plain_file) == "Hello, world!"

        bom_file = tmp_path / "bom.txt"
        bom_file.write_bytes(b"\xef\xbb\xbfBOM text")
        assert _read_utf8(bom_file) == "BOM text"

    def test_parse_text_file(self, tmp_path: Path) -> None:
        txt_path = tmp_path / "sample.txt"
        txt_path.write_text("Simple text line", encoding="utf-8")
        doc = parse_file(txt_path, "text/plain")
        assert doc.text == "Simple text line"
        assert doc.metadata == {"format": "txt"}
        assert doc.warnings == []

    def test_parse_markdown_file(self, tmp_path: Path) -> None:
        md_path = tmp_path / "sample.md"
        md_path.write_text("# Heading\n\n- Item 1\n- Item 2", encoding="utf-8")
        doc = parse_file(md_path, "text/markdown")
        assert "# Heading" in doc.text
        assert doc.metadata == {"format": "markdown"}
        assert doc.warnings == []


class TestJsonParser:
    """Test JSON file parsing, formatting, and error handling."""

    def test_parse_valid_json(self, tmp_path: Path) -> None:
        json_path = tmp_path / "data.json"
        raw_data = {"z_key": "last", "a_key": 123, "unicode": "Tiếng Việt"}
        json_path.write_text(json.dumps(raw_data), encoding="utf-8")

        doc = parse_file(json_path, "application/json")
        assert doc.metadata == {"format": "json"}
        assert doc.warnings == []
        # Formatted with indent 2 and sorted keys
        expected = json.dumps(raw_data, ensure_ascii=False, indent=2, sort_keys=True)
        assert doc.text == expected

    def test_parse_invalid_json_raises(self, tmp_path: Path) -> None:
        json_path = tmp_path / "broken.json"
        json_path.write_text("{not a valid json", encoding="utf-8")
        with pytest.raises(json.JSONDecodeError):
            parse_file(json_path, "application/json")


class TestCsvParser:
    """Test CSV file parsing including empty files, headers, and labeled columns."""

    def test_parse_empty_csv(self, tmp_path: Path) -> None:
        csv_path = tmp_path / "empty.csv"
        csv_path.write_text("", encoding="utf-8")
        assert _parse_csv(csv_path) == ""
        doc = parse_file(csv_path, "text/csv")
        assert doc.text == ""
        assert doc.metadata == {"format": "csv"}

    def test_parse_header_only_csv(self, tmp_path: Path) -> None:
        csv_path = tmp_path / "header.csv"
        csv_path.write_text("Name,Age,Role\n", encoding="utf-8")
        assert _parse_csv(csv_path) == "Name, Age, Role"

    def test_parse_regular_csv(self, tmp_path: Path) -> None:
        csv_path = tmp_path / "users.csv"
        csv_path.write_text("Name,Age\nAlice,30\nBob,25\n", encoding="utf-8")
        result = _parse_csv(csv_path)
        assert "Name: Alice\nAge: 30" in result
        assert "Name: Bob\nAge: 25" in result

    def test_parse_csv_with_extra_columns(self, tmp_path: Path) -> None:
        csv_path = tmp_path / "extra.csv"
        csv_path.write_text("Col1\nVal1,Val2,Val3\n", encoding="utf-8")
        result = _parse_csv(csv_path)
        assert "Col1: Val1" in result
        assert "column_2: Val2" in result
        assert "column_3: Val3" in result


class TestPdfParser:
    """Test PDF extraction, page limits, encryption, and empty text warnings."""

    def test_parse_pdf_without_extractable_text(self, tmp_path: Path) -> None:
        # Create a valid 1-page blank PDF
        pdf_path = tmp_path / "blank.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        with open(pdf_path, "wb") as f:
            writer.write(f)

        doc = parse_file(pdf_path, "application/pdf")
        assert doc.text == ""
        assert doc.metadata == {"format": "pdf", "pages": 1}
        assert doc.warnings == ["No extractable text; OCR is required"]

    def test_parse_pdf_with_text(self, tmp_path: Path) -> None:
        pdf_path = tmp_path / "mock.pdf"
        pdf_path.write_bytes(b"%PDF-1.4 dummy")

        mock_page = MagicMock()
        mock_page.extract_text.return_value = "Extracted PDF paragraph"
        mock_reader = MagicMock()
        mock_reader.is_encrypted = False
        mock_reader.pages = [mock_page]

        with patch("modules.ingestion.parsers.PdfReader", return_value=mock_reader):
            doc = parse_file(pdf_path, "application/pdf")
            assert doc.text == "Extracted PDF paragraph"
            assert doc.metadata == {"format": "pdf", "pages": 1}
            assert doc.warnings == []

    def test_parse_pdf_encrypted_raises(self, tmp_path: Path) -> None:
        pdf_path = tmp_path / "encrypted.pdf"
        pdf_path.write_bytes(b"%PDF-1.4 dummy")

        mock_reader = MagicMock()
        mock_reader.is_encrypted = True
        with patch("modules.ingestion.parsers.PdfReader", return_value=mock_reader):
            with pytest.raises(ValueError, match="Encrypted PDFs are not supported"):
                _parse_pdf(pdf_path, page_limit=PDF_PAGE_LIMIT)

    def test_parse_pdf_exceeds_page_limit_raises(self, tmp_path: Path) -> None:
        pdf_path = tmp_path / "large.pdf"
        pdf_path.write_bytes(b"%PDF-1.4 dummy")

        mock_reader = MagicMock()
        mock_reader.is_encrypted = False
        mock_reader.pages = [MagicMock() for _ in range(10)]

        with patch("modules.ingestion.parsers.PdfReader", return_value=mock_reader):
            with pytest.raises(ValueError, match="PDF exceeds the configured page limit"):
                _parse_pdf(pdf_path, page_limit=5)


class TestDocxParser:
    """Test DOCX paragraph and table extraction, archive checks, and size limits."""

    def test_parse_valid_docx(self, tmp_path: Path) -> None:
        docx_path = tmp_path / "doc.docx"
        doc = docx.Document()
        doc.add_paragraph("Paragraph 1")
        doc.add_paragraph("Paragraph 2")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "H1"
        table.rows[0].cells[1].text = "H2"
        table.rows[1].cells[0].text = "V1"
        table.rows[1].cells[1].text = "V2"
        doc.save(docx_path)

        parsed = parse_file(
            docx_path,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        assert "Paragraph 1" in parsed.text
        assert "Paragraph 2" in parsed.text
        assert "H1 | H2" in parsed.text
        assert "V1 | V2" in parsed.text
        assert parsed.metadata == {"format": "docx"}
        assert parsed.warnings == []

    def test_parse_invalid_docx_structure_raises(self, tmp_path: Path) -> None:
        bad_docx = tmp_path / "bad.docx"
        with zipfile.ZipFile(bad_docx, "w") as z:
            z.writestr("test.txt", "not a docx structure")

        with pytest.raises(ValueError, match="Invalid DOCX file"):
            _parse_docx(bad_docx, expanded_limit=DOCX_EXPANDED_LIMIT)

    def test_parse_docx_expanded_size_limit_raises(self, tmp_path: Path) -> None:
        docx_path = tmp_path / "oversized.docx"
        doc = docx.Document()
        doc.add_paragraph("Short text")
        doc.save(docx_path)

        # expanded_limit=1 byte to trigger limit
        with pytest.raises(ValueError, match="DOCX expanded size exceeds the configured limit"):
            _parse_docx(docx_path, expanded_limit=1)


class TestUnsupportedMimeAndBoundedParser:
    """Test unsupported MIME types and parse_file_bounded timeout/error handling."""

    def test_unsupported_mime_raises(self, tmp_path: Path) -> None:
        dummy_file = tmp_path / "image.png"
        dummy_file.write_bytes(b"\x89PNG")
        with pytest.raises(ValueError, match="Unsupported file type"):
            parse_file(dummy_file, "image/png")

    @pytest.mark.asyncio
    async def test_parse_file_bounded_success(self, tmp_path: Path) -> None:
        txt_path = tmp_path / "bounded.txt"
        txt_path.write_text("Bounded execution text", encoding="utf-8")

        result = await parse_file_bounded(
            path=txt_path,
            mime="text/plain",
            timeout_seconds=5,
            docx_expanded_limit=DOCX_EXPANDED_LIMIT,
            pdf_page_limit=PDF_PAGE_LIMIT,
        )
        assert isinstance(result, ParsedDocument)
        assert result.text == "Bounded execution text"

    @pytest.mark.asyncio
    async def test_parse_file_bounded_error_propagation(self, tmp_path: Path) -> None:
        txt_path = tmp_path / "dummy.unknown"
        txt_path.write_bytes(b"data")

        with pytest.raises(ValueError, match="Unsupported file type"):
            await parse_file_bounded(
                path=txt_path,
                mime="application/unknown",
                timeout_seconds=5,
                docx_expanded_limit=DOCX_EXPANDED_LIMIT,
                pdf_page_limit=PDF_PAGE_LIMIT,
            )
